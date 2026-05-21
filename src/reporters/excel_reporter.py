from __future__ import annotations
import io
from typing import Dict, List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

from ..comparators.base import DiffResult
from ..orchestrator import summarise

STATUS_FILLS = {
    "MATCH":         PatternFill("solid", fgColor="C6EFCE"),
    "MISMATCH":      PatternFill("solid", fgColor="FFEB9C"),
    "MISSING_IN_S4": PatternFill("solid", fgColor="FFC7CE"),
    "NEW_IN_S4":     PatternFill("solid", fgColor="BDD7EE"),
    "UNMAPPED_KEY":  PatternFill("solid", fgColor="E2EFDA"),
}
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _results_to_df(results: List[DiffResult], include_matches: bool) -> pd.DataFrame:
    rows = []
    for r in results:
        if not include_matches and r.status == "MATCH":
            continue
        if not r.field_diffs:
            rows.append({**r.primary_key, "STATUS": r.status, "FIELD": "", "R3_VALUE": "", "S4_VALUE": "", "COMP_TYPE": ""})
        else:
            for fd in r.field_diffs:
                rows.append({
                    **r.primary_key,
                    "STATUS": r.status,
                    "FIELD": fd.field_name,
                    "R3_VALUE": fd.r3_value,
                    "S4_VALUE": fd.s4_value,
                    "COMP_TYPE": fd.comparison_type,
                })
    return pd.DataFrame(rows)


def _write_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame):
    ws = wb.create_sheet(title=sheet_name)
    if df.empty:
        ws.append(["No records"])
        return

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        ws.append(row)
        for cell in ws[r_idx]:
            if r_idx == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
            else:
                status = df.iloc[r_idx - 2].get("STATUS", "")
                cell.fill = STATUS_FILLS.get(status, PatternFill())

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)


def generate_excel(results: Dict[str, List[DiffResult]], include_matches: bool = False) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    summary_df = summarise(results)
    _write_sheet(wb, "Summary", summary_df)

    for table, diffs in results.items():
        df = _results_to_df(diffs, include_matches)
        _write_sheet(wb, f"{table}_Diffs", df)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

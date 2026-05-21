import pytest
import pandas as pd
from src.comparators.base import DiffResult, FieldDiff
from src.reporters.excel_reporter import generate_excel, _results_to_df
from src.reporters.html_reporter import generate_html
from src.orchestrator import summarise


# ------------------------------------------------------------------ #
# Fixtures                                                             #
# ------------------------------------------------------------------ #

@pytest.fixture
def sample_results():
    return {
        "PLKO": [
            DiffResult("PLKO", {"PLNNR": "100", "PLNKN": "1", "WERKS": "1000"}, "MATCH"),
            DiffResult("PLKO", {"PLNNR": "200", "PLNKN": "1", "WERKS": "1000"}, "MISMATCH",
                       [FieldDiff("PLNBEZ", "direct", "OLD", "NEW")]),
            DiffResult("PLKO", {"PLNNR": "300", "PLNKN": "1", "WERKS": "1000"}, "MISSING_IN_S4"),
            DiffResult("PLKO", {"PLNNR": "400", "PLNKN": "1", "WERKS": "1000"}, "NEW_IN_S4"),
            DiffResult("PLKO", {"PLNNR": "500", "PLNKN": "1", "WERKS": "1000"}, "UNMAPPED_KEY"),
        ]
    }


# ------------------------------------------------------------------ #
# summarise                                                            #
# ------------------------------------------------------------------ #

def test_summarise_counts(sample_results):
    df = summarise(sample_results)
    row = df[df["Table"] == "PLKO"].iloc[0]
    assert row["MATCH"] == 1
    assert row["MISMATCH"] == 1
    assert row["MISSING_IN_S4"] == 1
    assert row["NEW_IN_S4"] == 1
    assert row["UNMAPPED_KEY"] == 1
    assert row["Total"] == 5


def test_summarise_match_rate(sample_results):
    df = summarise(sample_results)
    row = df[df["Table"] == "PLKO"].iloc[0]
    assert row["Match Rate %"] == 20.0


# ------------------------------------------------------------------ #
# _results_to_df                                                       #
# ------------------------------------------------------------------ #

def test_results_to_df_excludes_match_by_default(sample_results):
    df = _results_to_df(sample_results["PLKO"], include_matches=False)
    assert "MATCH" not in df["STATUS"].values


def test_results_to_df_includes_match_when_flag(sample_results):
    df = _results_to_df(sample_results["PLKO"], include_matches=True)
    assert "MATCH" in df["STATUS"].values


def test_results_to_df_mismatch_has_field_columns(sample_results):
    df = _results_to_df(sample_results["PLKO"], include_matches=False)
    mismatch_row = df[df["STATUS"] == "MISMATCH"].iloc[0]
    assert mismatch_row["FIELD"] == "PLNBEZ"
    assert mismatch_row["R3_VALUE"] == "OLD"
    assert mismatch_row["S4_VALUE"] == "NEW"


# ------------------------------------------------------------------ #
# Excel reporter                                                       #
# ------------------------------------------------------------------ #

def test_generate_excel_returns_bytes(sample_results):
    data = generate_excel(sample_results, include_matches=False)
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_generate_excel_valid_xlsx(sample_results):
    import io
    from openpyxl import load_workbook
    data = generate_excel(sample_results)
    wb = load_workbook(io.BytesIO(data))
    assert "Summary" in wb.sheetnames
    assert "PLKO_Diffs" in wb.sheetnames


def test_generate_excel_summary_sheet(sample_results):
    import io
    from openpyxl import load_workbook
    data = generate_excel(sample_results)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Summary"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert "Table" in headers
    assert "Match Rate %" in headers


# ------------------------------------------------------------------ #
# HTML reporter                                                        #
# ------------------------------------------------------------------ #

def test_generate_html_returns_string(sample_results):
    html = generate_html(sample_results, include_matches=False)
    assert isinstance(html, str)
    assert "<html" in html.lower()


def test_generate_html_contains_table_name(sample_results):
    html = generate_html(sample_results)
    assert "PLKO" in html


def test_generate_html_contains_status_filter(sample_results):
    html = generate_html(sample_results)
    assert "filterTable" in html


def test_generate_html_contains_summary_cards(sample_results):
    html = generate_html(sample_results)
    assert "Match Rate" in html or "match rate" in html


def test_generate_html_mismatch_count(sample_results):
    html = generate_html(sample_results)
    assert "MISMATCH" in html


def test_generate_html_self_contained(sample_results):
    html = generate_html(sample_results)
    assert "<script" in html
    assert "<style" in html
    assert "cdn" not in html.lower()
